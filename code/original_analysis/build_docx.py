import copy, html, json, re, sys, os
import docx
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.path.insert(0, ".")
import manuscript as M

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = f"{HERE}/../journal_submission/template_src/IJIES_Format.docx"
OUT = f"{HERE}/../journal_submission/Revised_Manuscript_12_Pages.docx"
FIGS = f"{HERE}/../journal_submission/figs"
COLW_TW = 4657  # one text column in twips
COLW_CM = 8.1

# --------------------------------------------------------------------------- references
CR = {c["ref"][0].lower().replace("'", "").replace(" ", ""): c["meta"] for c in json.load(open(f"{HERE}/refs_crossref.json"))}
CRL = json.load(open(f"{HERE}/refs_crossref.json"))


def initials(given):
    out = []
    for part in re.split(r"\s+", given.strip()):
        sub = [s for s in part.split("-") if s]
        out.append("-".join(s[0] + "." for s in sub))
    return " ".join(out)


def authors_str(au):
    names = [f"{initials(g)} {f}" for f, g in au]
    if len(names) == 1: return names[0]
    if len(names) == 2: return f"{names[0]} and {names[1]}"
    return ", ".join(names[:-1]) + ", and " + names[-1]


def jour(meta, title=None, extra_article=False):
    t = title or meta["title"]
    j = meta["journal"].replace("The Journal of Neuroscience", "Journal of Neuroscience")
    s = f"{authors_str(meta['authors'])}, “{t}”, <i>{j}</i>, Vol. {meta['vol']}"
    if meta.get("issue"): s += f", No. {meta['issue']}"
    pg = meta["page"]
    s += f", Art. No. {pg}" if re.match(r"^e?\d+$", pg) and not "-" in pg and (pg.startswith("e") or len(pg) >= 5) else f", pp. {pg}"
    return s + f", {meta['year']}."


# key -> (meta key in refs list order index)
KEYMAP = {"Blankertz": 0, "Bode": 1, "Boldt": 2, "Braun": 3, "Charles13": 4, "Charles14": 5, "Desender19a": 6, "Desender19b": 7, "Desender21": 8,
          "Gherman": 9, "Gluth": 10, "Herding": 11, "Kelly": 12, "OConnell": 13, "Philiastides": 14, "Ratcliff": 15, "Rausch": 16, "Sajda": 17,
          "Soon": 18, "Urai": 19, "vanVugt": 20}
TITLE_FIX = {"Bode": "Tracking the Unconscious Generation of Free Decisions Using Ultra-High Field fMRI"}
REFTXT = {}
for k, i in KEYMAP.items():
    m = CRL[i]["meta"]; REFTXT[k] = jour(m, TITLE_FIX.get(k))
REFTXT["Ledoit"] = "O. Ledoit and M. Wolf, “A well-conditioned estimator for large-dimensional covariance matrices”, <i>Journal of Multivariate Analysis</i>, Vol. 88, No. 2, pp. 365-411, 2004."
REFTXT["Liang"] = "K.-Y. Liang and S. L. Zeger, “Longitudinal data analysis using generalized linear models”, <i>Biometrika</i>, Vol. 73, No. 1, pp. 13-22, 1986."
REFTXT["Ojala"] = "M. Ojala and G. C. Garriga, “Permutation tests for studying classifier performance”, In: <i>Proc. of IEEE International Conf. on Data Mining</i>, pp. 908-913, 2009."
REFTXT["Combrisson"] = "E. Combrisson and K. Jerbi, “Exceeding chance level by chance: the caveat of theoretical chance levels in brain signal classification and statistical assessment of decoding accuracy”, <i>Journal of Neuroscience Methods</i>, Vol. 250, pp. 126-136, 2015."
REFTXT["Fleming"] = "S. M. Fleming and H. C. Lau, “How to measure metacognition”, <i>Frontiers in Human Neuroscience</i>, Vol. 8, Art. No. 443, 2014."
REFTXT["Dataset"] = "S. Gherman and M. G. Philiastides, “Simultaneous EEG-fMRI - Confidence in perceptual decisions”, <i>OpenNeuro</i>, ds002739, Version 1.0.0, DOI: 10.18112/openneuro.ds002739.v1.0.0, 2020."
REFTXT["Ojala"] = REFTXT["Ojala"]
# Fleming: Crossref gives article number 443 (no page); Dataset: dataset_description.json

# citation numbering by first appearance (document order)
ORDER = []
CITE = re.compile(r"\{([A-Za-z0-9_,]+)\}")


def scan(text):
    for m in CITE.finditer(text):
        for k in m.group(1).split(","):
            if k not in ORDER: ORDER.append(k)


for blk in M.BLOCKS:
    if blk[0] in ("p", "h1", "h2", "eq"): scan(blk[1])
    elif blk[0] == "fig": scan(blk[2])
    elif blk[0] == "tbl":
        scan(blk[1]); [scan(c) for r in blk[2] for c in r]
for k in ORDER: assert k in REFTXT, k
NUM = {k: i + 1 for i, k in enumerate(ORDER)}


def cite_repl(m):
    ks = sorted(NUM[k] for k in m.group(1).split(","))
    parts, i = [], 0
    while i < len(ks):
        j = i
        while j + 1 < len(ks) and ks[j + 1] == ks[j] + 1: j += 1
        parts.append(f"{ks[i]}" if i == j else f"{ks[i]}–{ks[j]}")
        i = j + 1
    return "[" + ", ".join(parts) + "]"


def render_cites(t): return CITE.sub(cite_repl, t)


# --------------------------------------------------------------------------- docx helpers
d = docx.Document(TPL)
body = d.element.body
P = list(d.paragraphs)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
proto = {"h1": copy.deepcopy(P[15]._p), "body": copy.deepcopy(P[16]._p), "cap": copy.deepcopy(P[42]._p), "chapter": copy.deepcopy(P[56]._p),
         "textbody": copy.deepcopy(P[57]._p), "refhead": copy.deepcopy(P[63]._p), "ref": copy.deepcopy(P[67]._p), "sh": copy.deepcopy(P[29]._p)}
tbl_style = d.tables[0].style
# 1) locate the section-break paragraphs and remove all template instructions after the title block
sect_paras = [p for p in body.iterchildren(qn("w:p")) if p.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None]
first_break = sect_paras[0]
final_sect = body.find(qn("w:sectPr"))
two_col = copy.deepcopy(sect_paras[1].find(qn("w:pPr") + "/" + qn("w:sectPr")))
start_removing = False
for el in list(body):
    if el is first_break: start_removing = True; continue
    if start_removing and el is not final_sect: body.remove(el)
two_col_end = copy.deepcopy(two_col)  # carried by a closing paragraph, as in the template; the body-level sectPr stays the template single-column one
two_col = final_sect  # insertion anchor: new content goes before the final body-level sectPr

for _p in list(body.iterchildren(qn("w:p"))):
    if _p.find(".//" + qn("w:drawing")) is not None:
        body.remove(_p); break
# 2) title block (keep template paragraphs; replace placeholder text; remove red colour)
def strip_red(p):
    for c in p.iter(qn("w:color")):
        if c.get(qn("w:val")) == "FF0000": c.set(qn("w:val"), "000000")


def set_runs(p, segments, strip_bi=False):
    """Keep first run's rPr as the base; segments: list of (text, dict(b,i,sup,sub))."""
    runs = p.findall(qn("w:r"))
    base = None
    for r in runs:
        if r.find(qn("w:rPr")) is not None: base = r.find(qn("w:rPr")); break
    base = copy.deepcopy(base) if base is not None else OxmlElement("w:rPr")
    for r in runs: p.remove(r)
    for text, f in segments:
        r = OxmlElement("w:r"); rp = copy.deepcopy(base)
        for tag in (("w:b", "w:i", "w:vertAlign", "w:color") if strip_bi else ("w:vertAlign", "w:color")):
            for e in rp.findall(qn(tag)): rp.remove(e)
        if f.get("b") and rp.find(qn("w:b")) is None: rp.append(OxmlElement("w:b"))
        if f.get("i") and rp.find(qn("w:i")) is None: rp.append(OxmlElement("w:i"))
        if f.get("sup") or f.get("sub"):
            va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript" if f.get("sup") else "subscript"); rp.append(va)
        # schema order: rFonts, b, i, ... , sz, ..., vertAlign -> reorder
        order = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath"]
        kids = sorted(list(rp), key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
        for e in list(rp): rp.remove(e)
        for e in kids: rp.append(e)
        r.append(rp); t = OxmlElement("w:t"); t.text = text; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); r.append(t); p.append(r)


def parse_markup(text):
    text = render_cites(text)
    toks = re.split(r"(</?(?:i|b|sup|sub)>)", text)
    st = {"i": 0, "b": 0, "sup": 0, "sub": 0}; out = []
    for tk in toks:
        m = re.match(r"<(/?)(i|b|sup|sub)>", tk)
        if m: st[m.group(2)] += -1 if m.group(1) else 1; continue
        if tk: out.append((html.unescape(tk), {k: v > 0 for k, v in st.items()}))
    return out


pp = lambda i: P[i]._p
title_p, auth_p, aff1, aff2, email_p, abs_p, kw_p = pp(1), pp(3), pp(5), pp(6), pp(7), pp(10), pp(11)
set_runs(title_p, [("Decoding Pre-Response EEG Signatures of Choice, Accuracy, and Confidence in Perceptual Decision-Making", {"b": True})], strip_bi=True)
set_runs(auth_p, [("[First Name Family Name]", {"b": True}), ("1*", {"b": True, "sup": True}), ("          ", {}), ("[First Name Family Name]", {"b": True}), ("2", {"b": True, "sup": True})], strip_bi=True)
set_runs(aff1, [("1", {"sup": True, "i": True}), ("[Full department and institution name, City, Country]", {"i": True})], strip_bi=True)
set_runs(aff2, [("2", {"sup": True, "i": True}), ("[Full department and institution name, City, Country]", {"i": True})], strip_bi=True)
set_runs(email_p, [("* Corresponding author’s Email: [author@institution.edu]", {})], strip_bi=True)
for e in (title_p, auth_p, aff1, aff2, email_p, abs_p, kw_p): strip_red(e)
set_runs(abs_p, [("Abstract: ", {"b": True}), (M.ABSTRACT, {})], strip_bi=True)
set_runs(kw_p, [("Keywords: ", {"b": True}), (M.KEYWORDS, {})], strip_bi=True)
# header placeholder dates in template stay as they are (publisher-filled)


def new_p(kind):
    p = copy.deepcopy(proto[kind])
    for r in p.findall(qn("w:r")): p.remove(r)
    for e in list(p):
        if e.tag not in (qn("w:pPr"),): p.remove(e)
    strip_red(p)
    body.insert(list(body).index(two_col), p)  # before final sectPr
    return p


def add_text(kind, text, pstyle_tweaks=None):
    p = new_p(kind)
    segs = parse_markup(text)
    # base run properties: take from prototype's first run
    tmp = copy.deepcopy(proto[kind]); r0 = tmp.find(qn("w:r"))
    base = copy.deepcopy(r0.find(qn("w:rPr"))) if r0 is not None and r0.find(qn("w:rPr")) is not None else OxmlElement("w:rPr")
    for e in list(base):
        if e.tag in (qn("w:color"),): base.remove(e)
    holder = OxmlElement("w:r"); holder.append(base); p.append(holder)
    set_runs(p, segs)
    if pstyle_tweaks: pstyle_tweaks(p)
    return p


def set_spacing(p, before=None, after=None, line=None, rule="auto", keep_next=False):
    pPr = p.find(qn("w:pPr"))
    sp = pPr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing"); pPr.append(sp)
    for k, v in (("before", before), ("after", after), ("line", line)):
        if v is not None: sp.set(qn("w:" + k), str(v))
    for a in ("beforeLines", "afterLines"):
        if sp.get(qn("w:" + a)) is not None and (before is not None or after is not None): del sp.attrib[qn("w:" + a)]
    if line is not None: sp.set(qn("w:lineRule"), rule)
    if keep_next and pPr.find(qn("w:keepNext")) is None:
        kn = OxmlElement("w:keepNext"); pPr.insert(1 if pPr.find(qn("w:pStyle")) is not None else 0, kn)
    # keep pPr children in schema order
    order = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "widowControl", "numPr", "tabs", "autoSpaceDE", "autoSpaceDN", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "jc", "outlineLvl", "rPr", "sectPr"]
    kids = sorted(list(pPr), key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 98)
    for e in list(pPr): pPr.remove(e)
    for e in kids: pPr.append(e)


def h1(t): p = add_text("h1", t); set_spacing(p, before=160, after=60, keep_next=True); return p


def h2(t):
    p = add_text("sh", t); pPr = p.find(qn("w:pPr"))
    for e in pPr.findall(qn("w:numPr")): pPr.remove(e)
    ind = pPr.find(qn("w:ind"))
    if ind is not None: pPr.remove(ind)
    for e in pPr.findall(qn("w:tabs")): pPr.remove(e)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "0"); ind.set(qn("w:firstLine"), "0"); pPr.append(ind)
    jc = pPr.find(qn("w:jc"))
    if jc is None: jc = OxmlElement("w:jc"); pPr.append(jc)
    jc.set(qn("w:val"), "left")
    set_spacing(p, before=100, after=40, keep_next=True)
    return p


def para(t): p = add_text("body", t); set_spacing(p, before=0, after=0, line=240); return p


H1N = [0]; H2N = [0]
_h1, _h2 = h1, h2


def numbered_h2(t):
    return _h2(f"{H1N[0]}.{H2N[0]} {t}")


def numbered_h1(t):
    return _h1(t)


def equation(markup, num):
    p = add_text("body", markup)
    pPr = p.find(qn("w:pPr"))
    for e in pPr.findall(qn("w:ind")): pPr.remove(e)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "283"); ind.set(qn("w:firstLine"), "0"); pPr.append(ind)
    tabs = OxmlElement("w:tabs"); tb = OxmlElement("w:tab"); tb.set(qn("w:val"), "right"); tb.set(qn("w:pos"), str(COLW_TW)); tabs.append(tb); pPr.append(tabs)
    jc = pPr.find(qn("w:jc")); jc.set(qn("w:val"), "left")
    # equation, then right tab and number
    r = OxmlElement("w:r"); rp = OxmlElement("w:rPr"); rf = OxmlElement("w:rFonts")
    for a in ("ascii", "hAnsi", "cs"): rf.set(qn("w:" + a), "Times New Roman")
    rp.append(rf); r.append(rp); tab = OxmlElement("w:tab"); r.append(tab); p.append(r)
    r2 = OxmlElement("w:r"); r2.append(copy.deepcopy(rp)); t = OxmlElement("w:t"); t.text = f"({num})"; r2.append(t); p.append(r2)
    set_spacing(p, before=120, after=120, line=240)
    return p


def figure(fn, caption, fig_no):
    p = new_p("body"); pPr = p.find(qn("w:pPr"))
    for e in pPr.findall(qn("w:ind")): pPr.remove(e)
    jc = pPr.find(qn("w:jc")); jc.set(qn("w:val"), "center")
    set_spacing(p, before=80, after=20, line=240, keep_next=True)
    para_obj = docx.text.paragraph.Paragraph(p, d._body)
    run = para_obj.add_run(); run.add_picture(f"{FIGS}/{fn}", width=Cm(COLW_CM))
    c = add_text("cap", f"Fig. {fig_no}. {caption}")
    set_spacing(c, before=0, after=100, line=240)
    pPr = c.find(qn("w:pPr")); jc = pPr.find(qn("w:jc")); jc.set(qn("w:val"), "center")
    return p


def table(title, rows, widths_cm, tbl_no, header=True, caption=None, left=False):
    title = re.sub(r"^Table \d+\. ", "", title)
    cp = add_text("cap", caption or f"Table {tbl_no}. {title}")
    set_spacing(cp, before=80, after=40, line=240, keep_next=True)
    pPr = cp.find(qn("w:pPr")); pPr.find(qn("w:jc")).set(qn("w:val"), "center")
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = tbl_style
    el = t._tbl; body.remove(el); body.insert(list(body).index(two_col), el)
    tw = sum(widths_cm)
    tblPr = el.tblPr
    for e in tblPr.findall(qn("w:tblW")): tblPr.remove(e)
    tw_e = OxmlElement("w:tblW"); tw_e.set(qn("w:w"), str(int(tw / 2.54 * 1440))); tw_e.set(qn("w:type"), "dxa"); tblPr.append(tw_e)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); tblPr.append(jc)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("left", 40), ("right", 40)):
        e = OxmlElement("w:" + side); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
    order = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
    kids = sorted(list(tblPr), key=lambda e: order.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in order else 99)
    for e in list(tblPr): tblPr.remove(e)
    for e in kids: tblPr.append(e)
    grid = el.find(qn("w:tblGrid"))
    for g, wcm in zip(grid.findall(qn("w:gridCol")), widths_cm): g.set(qn("w:w"), str(int(wcm / 2.54 * 1440)))
    for ri, row in enumerate(rows):
        trPr = t.rows[ri]._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        for ci, txt in enumerate(row):
            cell = t.cell(ri, ci); cell.width = Cm(widths_cm[ci])
            pc = cell.paragraphs[0]._p
            pPr = pc.get_or_add_pPr()
            if ri < len(rows) - 1: pPr.append(OxmlElement("w:keepNext"))
            sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "0"); sp.set(qn("w:after"), "0"); sp.set(qn("w:line"), "240"); sp.set(qn("w:lineRule"), "auto"); pPr.append(sp)
            jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "left" if (ci == 0 or not header or left) else "center"); pPr.append(jc)
            segs = parse_markup(txt)
            for text, f in segs:
                r = OxmlElement("w:r"); rp = OxmlElement("w:rPr"); rf = OxmlElement("w:rFonts")
                for a in ("ascii", "hAnsi", "cs"): rf.set(qn("w:" + a), "Times New Roman")
                rp.append(rf)
                if (header and ri == 0) or f.get("b"): rp.append(OxmlElement("w:b"))
                if f.get("i"): rp.append(OxmlElement("w:i"))
                sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "20"); rp.append(sz); szc = OxmlElement("w:szCs"); szc.set(qn("w:val"), "20"); rp.append(szc)
                r.append(rp); te = OxmlElement("w:t"); te.text = text; te.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); r.append(te); pc.append(r)
    after = add_text("body", " "); set_spacing(after, before=0, after=0, line=120, rule="exact")  # small gap after table
    return el


def algorithm(title, lines):
    return table(title, [[l] for l in lines], [COLW_CM], 0, header=False, caption=title)


# --------------------------------------------------------------------------- assemble body
fig_no = [0]; tbl_no = [0]
for blk in M.BLOCKS:
    k = blk[0]
    if k == "h1": H1N[0] += 1; H2N[0] = 0; numbered_h1(blk[1])
    elif k == "h2": H2N[0] += 1; numbered_h2(blk[1])
    elif k == "p": para(blk[1])
    elif k == "eq": equation(blk[1], blk[2])
    elif k == "fig": fig_no[0] += 1; figure(blk[1], blk[2], fig_no[0])
    elif k == "tbl": tbl_no[0] += 1; table(blk[1], blk[2], blk[3], tbl_no[0], left=(len(blk) > 4 and blk[4] == "left"))
    elif k == "alg": algorithm(blk[1], blk[2])

# declarations
for title, txt in M.DECL:
    c = add_text("chapter", title)
    tb = add_text("textbody", txt)
    for e in list(tb.iter(qn("w:sz"))): e.set(qn("w:val"), "22")
# references
add_text("refhead", "References")
for k in ORDER:
    p = add_text("ref", REFTXT[k])
    for e in p.iter(qn("w:color")): e.set(qn("w:val"), "000000")
    ppr = p.find(qn("w:pPr"))
    for e in ppr.findall(qn("w:tabs")): ppr.remove(e)
    ind = ppr.find(qn("w:ind")); ind.set(qn("w:left"), "480"); ind.set(qn("w:hanging"), "480")
endp = OxmlElement("w:p"); eppr = OxmlElement("w:pPr"); esp = OxmlElement("w:spacing"); esp.set(qn("w:before"), "0"); esp.set(qn("w:after"), "0"); esp.set(qn("w:line"), "20"); esp.set(qn("w:lineRule"), "exact")
eppr.append(esp); eppr.append(two_col_end); endp.append(eppr); body.insert(list(body).index(final_sect), endp)
for el in body.iter():
    for att in list(el.attrib):
        if att.endswith("}paraId") or att.endswith("}textId"): del el.attrib[att]
# plain document: no logo banner, no header, no footer
for sec in d.sections:
    for part in (sec.header, sec.footer):
        part.is_linked_to_previous = False
        el = part._element
        for child in list(el): el.remove(child)
        el.append(OxmlElement("w:p"))
d.core_properties.title = "Decoding Pre-Response EEG Signatures of Choice, Accuracy, and Confidence in Perceptual Decision-Making"
d.core_properties.author = ""
d.save(OUT)
json.dump({"order": ORDER, "numbers": NUM, "reftext": {k: REFTXT[k] for k in ORDER}}, open(f"{HERE}/ref_order.json", "w"), indent=1)
print("saved", OUT, "figs", fig_no[0], "tables", tbl_no[0], "refs", len(ORDER))
